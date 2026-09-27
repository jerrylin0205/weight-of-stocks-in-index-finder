from __future__ import annotations

import io
import re
from datetime import date

from dateutil import parser as dateparser
from openpyxl import load_workbook

from ..models import Holding, HoldingsSnapshot
from .base import (
    FetchContext,
    FetchError,
    Source,
    clean_text,
    http_get,
    parse_number,
    save_raw,
)
from ..normalize import normalize_ticker

URL = (
    "https://www.ssga.com/us/en/intermediary/library-content/products/"
    "fund-data/etfs/us/holdings-daily-us-en-{t}.xlsx"
)

_AS_OF = re.compile(
    r"as of\s+(\d{1,2}[-/][A-Za-z]{3}[-/]\d{2,4}|\d{4}-\d{2}-\d{2}|[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4})",
    re.I,
)


class SSGASource(Source):
    key = "ssga"

    def fetch(self, cfg, ctx: FetchContext) -> HoldingsSnapshot:
        t = str(cfg.params.get("ssga_ticker", cfg.ticker)).lower()
        r = http_get(URL.format(t=t), expect=lambda r: r.content[:2] == b"PK")
        as_of, holdings = _parse(r.content)
        raw = save_raw(ctx.raw_dir, cfg.ticker, r.content, "xlsx", as_of)
        return HoldingsSnapshot(cfg.ticker, as_of, self.key, holdings, raw)


def _col(header: list[str], *names: str) -> int | None:
    for want in names:
        for i, h in enumerate(header):
            if h == want:
                return i
    for want in names:
        for i, h in enumerate(header):
            if want in h:
                return i
    return None


def _parse(blob: bytes) -> tuple[date, list[Holding]]:
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    ws = wb.active
    rows = [[c.value for c in row] for row in ws.iter_rows()]

    as_of: date | None = None
    header_idx: int | None = None
    header: list[str] = []
    for i, row in enumerate(rows):
        joined = " ".join(str(c) for c in row if c is not None)
        if as_of is None:
            m = _AS_OF.search(joined)
            if m:
                try:
                    as_of = dateparser.parse(m.group(1)).date()
                except (ValueError, OverflowError):
                    pass
        low = [str(c).strip().lower() if c is not None else "" for c in row]
        if "ticker" in low and any("weight" in c for c in low):
            header_idx, header = i, low
            break

    if header_idx is None:
        raise FetchError("SSGA XLSX: 找不到含 Ticker / Weight 的標題列")

    ti = _col(header, "ticker")
    wi = _col(header, "weight")
    ni = _col(header, "name")
    holdings: list[Holding] = []
    for row in rows[header_idx + 1 :]:
        if not row or all(c is None for c in row):
            continue
        tk = clean_text(row[ti]) if ti is not None and ti < len(row) else ""
        nm = clean_text(row[ni]) if ni is not None and ni < len(row) else ""
        wt = parse_number(row[wi]) if wi is not None and wi < len(row) else None
        if not tk or wt is None or wt <= 0:
            continue
        low_tk, low_nm = tk.lower(), nm.lower()
        if low_tk in ("-", "cash", "usd", "cash&other") or "unrealized" in low_nm:
            continue
        if "cash" in low_nm and not tk.isalpha():
            continue
        holdings.append(Holding(normalize_ticker(tk, "US"), tk, nm, wt))

    return as_of or date.today(), holdings
